"""
Testy jednostkowe dla webapp/reports.py (eksport PDF i Excel).
"""

import io

import openpyxl
import pytest

from mcdm.controller import MCDMController
from mcdm.models.decision_problem import DecisionProblem
from mcdm.strategies import AhpStrategy, ElectreIStrategy, PrometheeStrategy, TopsisStrategy
from webapp.reports import build_excel_report, build_pdf_report

EXAMPLES_DIR = "data/examples"


def _make_results():
    problem = DecisionProblem.from_json(
        f"{EXAMPLES_DIR}/sewage_network_variants.json"
    )
    controller = MCDMController()
    controller.register_strategy(TopsisStrategy())
    controller.register_strategy(AhpStrategy())
    controller.register_strategy(PrometheeStrategy())
    controller.register_strategy(ElectreIStrategy())
    return problem, controller.run_all(problem)


def test_build_pdf_report_returns_valid_pdf_bytes():
    problem, results = _make_results()
    pdf_bytes = build_pdf_report(problem, results, "Test source")

    assert isinstance(pdf_bytes, bytes)
    assert pdf_bytes[:4] == b"%PDF"
    assert len(pdf_bytes) > 1000


def test_build_pdf_report_works_without_gaia_for_two_criteria():
    """Dla problemu z <3 kryteriami PDF nadal powinien sie zbudowac
    (bez sekcji GAIA), bez rzucania wyjatku."""
    problem = DecisionProblem(
        matrix=[[1, 7], [2, 4], [3, 1]],
        weights=[0.5, 0.5],
        directions=["max", "max"],
        alternative_names=["Alt1", "Alt2", "Alt3"],
        criterion_names=["K1", "K2"],
    )
    controller = MCDMController()
    controller.register_strategy(TopsisStrategy())
    results = controller.run_all(problem)

    pdf_bytes = build_pdf_report(problem, results, "Test 2 kryteria")
    assert pdf_bytes[:4] == b"%PDF"


def test_build_excel_report_returns_valid_workbook():
    problem, results = _make_results()
    excel_bytes = build_excel_report(problem, results, "Test source")

    assert isinstance(excel_bytes, bytes)
    assert excel_bytes[:2] == b"PK"  # xlsx = zip archive

    workbook = openpyxl.load_workbook(io.BytesIO(excel_bytes))
    expected_sheets = {"Macierz decyzyjna", "TOPSIS", "AHP", "PROMETHEE",
                        "ELECTRE I", "Zgodnosc rankingow"}
    assert expected_sheets.issubset(set(workbook.sheetnames))


def test_build_excel_report_ranking_sheet_matches_result():
    problem, results = _make_results()
    excel_bytes = build_excel_report(problem, results, "Test source")
    workbook = openpyxl.load_workbook(io.BytesIO(excel_bytes))

    sheet = workbook["TOPSIS"]
    rows = list(sheet.iter_rows(values_only=True))
    assert rows[0] == ("Pozycja", "Alternatywa", "Wynik")

    topsis_result = results["TOPSIS"]
    expected_leader = topsis_result.as_ordered_names()[0]
    assert rows[1][1] == expected_leader
    assert rows[1][0] == 1
