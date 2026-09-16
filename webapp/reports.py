"""
Generowanie raportow do pobrania z widoku wynikow (PU4) -- konkretny,
namacalny artefakt do zalacznika pracy inzynierskiej.

- build_pdf_report()   -- PDF z macierza decyzyjna, wykresami
  (radar, GAIA) i tabelami rankingow wszystkich metod (reportlab).
- build_excel_report() -- skoroszyt Excel z osobnym arkuszem na
  macierz decyzyjna, kazda metode oraz macierz zgodnosci rankingow
  (openpyxl).

Oba generatory przyjmuja te same dane co widok results.html (problem,
slownik wynikow z MCDMController.run_all()), wiec raport zawsze
odzwierciedla dokladnie to, co widac na ekranie.
"""

from __future__ import annotations

import io
from datetime import datetime

import openpyxl
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import (
    Image,
    KeepTogether,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

from mcdm.evaluation.rank_correlation import compare_all_pairs
from mcdm.models.decision_problem import DecisionProblem
from mcdm.strategies.base import RankingResult
from mcdm.visualization._utils import figure_to_png_bytes
from mcdm.visualization.gaia import plot_gaia_plane
from mcdm.visualization.radar import plot_radar_chart

BRAND_COLOR = colors.HexColor("#4c5fe0")
ROW_ALT_COLOR = colors.HexColor("#f5f6fb")
GRID_COLOR = colors.HexColor("#e2e5f1")


def build_pdf_report(
    problem: DecisionProblem,
    results: dict[str, RankingResult],
    source_label: str,
) -> bytes:
    """
    Buduje raport PDF: naglowek, macierz decyzyjna, wykres radarowy,
    plaszczyzna GAIA (jesli problem ma >=3 kryteria), rankingi
    wszystkich metod oraz macierz zgodnosci Kendall/Spearman.

    Returns
    -------
    bytes -- gotowa zawartosc pliku PDF.
    """
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        leftMargin=2 * cm,
        rightMargin=2 * cm,
        topMargin=2 * cm,
        bottomMargin=2 * cm,
    )
    styles = getSampleStyleSheet()
    muted_style = ParagraphStyle(
        "muted", parent=styles["Normal"], textColor=colors.grey, fontSize=9
    )

    story = []

    story.append(Paragraph("Raport MCDM Toolkit", styles["Title"]))
    story.append(Paragraph(f"Zrodlo problemu decyzyjnego: {source_label}", styles["Normal"]))
    story.append(
        Paragraph(
            f"Wygenerowano: {datetime.now().strftime('%Y-%m-%d %H:%M')}", muted_style
        )
    )
    story.append(Spacer(1, 0.5 * cm))

    story.append(Paragraph("Macierz decyzyjna", styles["Heading2"]))
    matrix_rows = [["Alternatywa"] + problem.criterion_names]
    for name, row in zip(problem.active_names, problem.active_matrix.tolist()):
        matrix_rows.append([name] + [f"{v:.2f}" for v in row])
    story.append(Table(matrix_rows, hAlign="LEFT", style=_table_style(), repeatRows=1))
    story.append(Spacer(1, 0.5 * cm))

    story.append(Paragraph("Profile wydajnosci alternatyw", styles["Heading2"]))
    radar_bytes = figure_to_png_bytes(plot_radar_chart(problem))
    story.append(
        Image(io.BytesIO(radar_bytes), width=14 * cm, height=14 * cm, kind="proportional")
    )

    if problem.n_criteria >= 3:
        story.append(Spacer(1, 0.5 * cm))
        story.append(Paragraph("Plaszczyzna GAIA", styles["Heading2"]))
        gaia_bytes = figure_to_png_bytes(plot_gaia_plane(problem))
        story.append(
            Image(io.BytesIO(gaia_bytes), width=15 * cm, height=11 * cm, kind="proportional")
        )

    story.append(PageBreak())

    for method_name, result in results.items():
        rank_rows = [["Pozycja", "Alternatywa", "Wynik"]]
        for pos, idx in enumerate(result.ranking, start=1):
            rank_rows.append(
                [str(pos), result.alternative_names[idx], f"{result.scores[idx]:.4f}"]
            )
        story.append(
            KeepTogether(
                [
                    Paragraph(f"Ranking: {method_name}", styles["Heading2"]),
                    Table(rank_rows, hAlign="LEFT", style=_table_style(), repeatRows=1),
                ]
            )
        )
        story.append(Spacer(1, 0.4 * cm))

    correlations = compare_all_pairs(results)
    corr_rows = [["Metoda A", "Metoda B", "Kendall tau", "Spearman rho"]]
    for (a, b), corr in correlations.items():
        corr_rows.append(
            [a, b, f"{corr['kendall_tau']:.3f}", f"{corr['spearman_rho']:.3f}"]
        )
    story.append(
        KeepTogether(
            [
                Paragraph("Zgodnosc rankingow (Kendall tau / Spearman rho)", styles["Heading2"]),
                Table(corr_rows, hAlign="LEFT", style=_table_style(), repeatRows=1),
            ]
        )
    )

    doc.build(story)
    buffer.seek(0)
    return buffer.getvalue()


def _table_style() -> TableStyle:
    return TableStyle(
        [
            ("BACKGROUND", (0, 0), (-1, 0), BRAND_COLOR),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("GRID", (0, 0), (-1, -1), 0.5, GRID_COLOR),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, ROW_ALT_COLOR]),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("TOPPADDING", (0, 0), (-1, -1), 4),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ]
    )


def build_excel_report(
    problem: DecisionProblem,
    results: dict[str, RankingResult],
    source_label: str,
) -> bytes:
    """
    Buduje skoroszyt Excel: arkusz z macierza decyzyjna, po jednym
    arkuszu na kazda metode MCDM oraz arkusz zgodnosci rankingow.

    Returns
    -------
    bytes -- gotowa zawartosc pliku .xlsx.
    """
    workbook = openpyxl.Workbook()
    header_fill = PatternFill(start_color="4C5FE0", end_color="4C5FE0", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)

    sheet = workbook.active
    sheet.title = "Macierz decyzyjna"
    sheet.append(["Zrodlo:", source_label])
    sheet.append([])
    sheet.append(["Alternatywa"] + problem.criterion_names)
    for name, row in zip(problem.active_names, problem.active_matrix.tolist()):
        sheet.append([name] + [float(v) for v in row])
    _style_header_row(sheet, header_fill, header_font, header_row=3)
    _autofit_columns(sheet)

    for method_name, result in results.items():
        sheet = workbook.create_sheet(title=_safe_sheet_title(method_name))
        sheet.append(["Pozycja", "Alternatywa", "Wynik"])
        for pos, idx in enumerate(result.ranking, start=1):
            sheet.append([pos, result.alternative_names[idx], round(float(result.scores[idx]), 4)])
        _style_header_row(sheet, header_fill, header_font, header_row=1)
        _autofit_columns(sheet)

    sheet = workbook.create_sheet(title="Zgodnosc rankingow")
    sheet.append(["Metoda A", "Metoda B", "Kendall tau", "Spearman rho"])
    correlations = compare_all_pairs(results)
    for (a, b), corr in correlations.items():
        sheet.append([a, b, round(float(corr["kendall_tau"]), 3), round(float(corr["spearman_rho"]), 3)])
    _style_header_row(sheet, header_fill, header_font, header_row=1)
    _autofit_columns(sheet)

    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def _safe_sheet_title(name: str) -> str:
    """Nazwy arkuszy Excela sa ograniczone do 31 znakow i nie moga zawierac
    niektorych znakow specjalnych (np. '/') -- ELECTRE I ma spacje, co jest OK."""
    return name.replace("/", "-")[:31]


def _style_header_row(sheet, fill, font, header_row: int) -> None:
    for cell in sheet[header_row]:
        cell.fill = fill
        cell.font = font


def _autofit_columns(sheet) -> None:
    for column_cells in sheet.columns:
        length = max(
            (len(str(cell.value)) for cell in column_cells if cell.value is not None),
            default=0,
        )
        sheet.column_dimensions[column_cells[0].column_letter].width = min(
            max(length + 2, 10), 40
        )
